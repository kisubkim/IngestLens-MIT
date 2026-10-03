"""Office files without LibreOffice, plus structure hints for either path.

render_native(): python-docx / python-pptx / openpyxl -> one HTML block per unit (document, slide,
sheet part) -> PDF with PyMuPDF Story. Layout is approximated, but reading order, heading sizes,
tables and images survive, which is what profiling, parsing and chunking need.

extract_hints(): things the PDF loses even with LibreOffice: speaker notes, slide titles and the
data behind native PowerPoint charts.
"""

import html
import re
from pathlib import Path

import pymupdf

A4 = pymupdf.paper_rect("a4")
SLIDE = pymupdf.Rect(0, 0, 960, 540)  # 16:9 in points
MARGIN = 36

BASE_CSS = """
* { font-family: kr; }
body { font-size: 10pt; line-height: 1.45; }
h1 { font-size: 20pt; margin: 0 0 8pt; }
h2 { font-size: 16pt; margin: 10pt 0 6pt; }
h3 { font-size: 13pt; margin: 8pt 0 4pt; }
p { margin: 0 0 5pt; }
table { border-collapse: collapse; margin: 6pt 0; }
td, th { border: 0.6pt solid #555; padding: 2pt 4pt; font-size: 9pt; vertical-align: top; }
th { font-weight: bold; }
.caption { font-size: 9pt; }
.slide { font-size: 18pt; }
.slide h1 { font-size: 30pt; margin-bottom: 14pt; }
.slide td, .slide th, .slide .caption { font-size: 14pt; }
"""


def _e(s) -> str:
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    return html.escape("" if s is None else str(s))


def _table_html(rows: list[list], header: bool = True, caption: str | None = None) -> str:
    out = ["<table>"]
    for i, row in enumerate(rows):
        tag = "th" if header and i == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{_e(c)}</{tag}>" for c in row) + "</tr>")
    out.append("</table>")
    cap = f'<p class="caption">{_e(caption)}</p>' if caption else ""
    return cap + "".join(out)


# ---------- docx ----------

def _docx_units(path: Path, archive: pymupdf.Archive) -> tuple[list[str], dict]:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    parts, headings, n_img = [], [], 0
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, d)
            style = (p.style.name if p.style is not None else "") or ""
            text = p.text.strip()
            m = re.match(r"(?:Heading|제목)\s*(\d)", style)
            level = 1 if style in ("Title", "제목") else int(m.group(1)) if m else 0
            for blip in el.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                if rid and rid in d.part.related_parts:
                    n_img += 1
                    name = f"docx_img_{n_img}"
                    archive.add(d.part.related_parts[rid].blob, name)
                    parts.append(f'<p><img src="{name}" style="max-width:100%"/></p>')
            if not text:
                continue
            if level:
                h = min(level, 3)
                headings.append({"level": level, "text": text})
                parts.append(f"<h{h}>{_e(text)}</h{h}>")
            else:
                parts.append(f"<p>{_e(text)}</p>")
        elif el.tag == qn("w:tbl"):
            t = Table(el, d)
            parts.append(_table_html([[c.text.strip() for c in r.cells] for r in t.rows]))
    return ["".join(parts)], {"kind": "docx", "headings": headings, "images": n_img}


# ---------- pptx ----------

def _chart_rows(chart) -> tuple[str, list[list]]:
    title = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else ""
    plot = chart.plots[0]
    cats = [str(c) for c in plot.categories]
    series = list(plot.series)
    rows = [["항목"] + [s.name for s in series]]
    for i, c in enumerate(cats):
        rows.append([c] + [s.values[i] if i < len(s.values) else "" for s in series])
    return title, rows


def _shapes(shapes):
    """Flatten group shapes, in visual order (top, then left)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    flat = []
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(_shapes(s.shapes))
        else:
            flat.append(s)
    return sorted(flat, key=lambda s: ((s.top or 0), (s.left or 0)))


def _pptx_slides(path: Path, archive: pymupdf.Archive | None) -> tuple[list[str], dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    units, slides = [], []
    for i, slide in enumerate(prs.slides):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        body, charts = [], []
        if title:
            body.append(f"<h1>{_e(title)}</h1>")
        for s in _shapes(slide.shapes):
            if title_shape is not None and s.shape_id == title_shape.shape_id:
                continue
            if getattr(s, "has_chart", False) and s.has_chart:
                ctitle, rows = _chart_rows(s.chart)
                charts.append({"title": ctitle, "rows": rows})
                body.append(_table_html(rows, caption=f"차트: {ctitle}" if ctitle else "차트"))
            elif getattr(s, "has_table", False) and s.has_table:
                body.append(_table_html([[c.text.strip() for c in r.cells] for r in s.table.rows]))
            elif s.shape_type == MSO_SHAPE_TYPE.PICTURE and archive is not None:
                name = f"slide{i}_img{s.shape_id}"
                archive.add(s.image.blob, name)
                body.append(f'<p><img src="{name}" style="max-height:260pt"/></p>')
            elif s.has_text_frame:
                for para in s.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        indent = "&#160;&#160;&#160;&#160;" * (para.level or 0)
                        body.append(f"<p>{indent}{_e(t)}</p>")
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        units.append('<div class="slide">' + ("".join(body) or "<p></p>") + "</div>")
        slides.append({"slide": i, "title": title, "notes": notes, "charts": charts})
    return units, {"kind": "pptx", "slides": slides}


# ---------- xlsx ----------

def _xlsx_units(path: Path, rows_per_page: int, max_rows: int) -> tuple[list[str], dict]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    units, sheets = [], []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if len(rows) >= max_rows + 1:
                break
            rows.append(["" if v is None else v for v in r])
        while rows and not any(str(v).strip() for v in rows[-1]):
            rows.pop()
        width = max((i + 1 for r in rows for i, v in enumerate(r) if str(v).strip()), default=0)
        rows = [r[:width] for r in rows]
        truncated = ws.max_row is not None and ws.max_row > max_rows + 1
        sheets.append({"sheet": ws.title, "rows": max(0, len(rows) - 1), "cols": width, "truncated": truncated})
        if not rows:
            continue
        header, data = rows[0], rows[1:] or [[]]
        # One unit per block of rows, header repeated, so every page (and chunk) is a self-describing table.
        for start in range(0, len(data), rows_per_page):
            block = data[start:start + rows_per_page]
            label = f"{ws.title} (행 {start + 2}–{start + 1 + len(block)})"
            units.append(f"<h2>{_e(label)}</h2>" + _table_html([header] + block))
    wb.close()
    return units, {"kind": "xlsx", "sheets": sheets}


# ---------- render ----------

def _write(units: list[str], out: Path, page: pymupdf.Rect, archive: pymupdf.Archive) -> list[int]:
    """Each unit starts on a new page and may flow onto more. Returns the first page index of every unit."""
    # PyMuPDF's built-in CJK font: no system font or extra package needed offline.
    archive.add(pymupdf.Font("korea").buffer, "kr.ttf")
    css = '@font-face { font-family: kr; src: url(kr.ttf); }' + BASE_CSS
    writer = pymupdf.DocumentWriter(str(out))
    where = page + (MARGIN, MARGIN, -MARGIN, -MARGIN)
    starts, n = [], 0
    for u in units:
        starts.append(n)
        story = pymupdf.Story(html=u, user_css=css, archive=archive)
        more = True
        while more:
            dev = writer.begin_page(page)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
            n += 1
    writer.close()
    return starts


def render_native(src: Path, out_dir: Path, fmt: str, xlsx_rows_per_page: int = 24, xlsx_max_rows: int = 5000) -> tuple[Path, dict]:
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf = out_dir / (src.stem + ".native.pdf")
    archive = pymupdf.Archive()
    if fmt == "docx":
        units, hints = _docx_units(src, archive)
        page = A4
    elif fmt == "pptx":
        units, hints = _pptx_slides(src, archive)
        page = SLIDE
    elif fmt == "xlsx":
        units, hints = _xlsx_units(src, xlsx_rows_per_page, xlsx_max_rows)
        page = pymupdf.paper_rect("a4-l")
    else:
        raise ValueError(f"no native renderer for .{fmt}")
    starts = _write(units or ["<p></p>"], pdf, page, archive)
    if hints["kind"] == "pptx":
        for s, first in zip(hints["slides"], starts):
            s["page"] = first
    hints["unit_pages"] = starts
    return pdf, hints


def extract_hints(src: Path, fmt: str) -> dict:
    """Hints for a LibreOffice-converted file (LibreOffice renders one PDF page per slide)."""
    if fmt == "pptx":
        _, hints = _pptx_slides(src, None)
        for s in hints["slides"]:
            s["page"] = s["slide"]
        return hints
    if fmt == "docx":
        _, hints = _docx_units(src, pymupdf.Archive())
        return hints
    return {}
